"""CSV / XLSX export of suggestions and a copy-ready WhatsApp table."""

from __future__ import annotations

import csv
import io

from .util.money import format_zar
from .util.timeutil import fmt_local

COLUMNS = ["Tier", "Score", "Site", "Auction", "Lot", "Ring", "Sex", "Seller", "Current", "Your max",
           "Closes", "Match", "Families", "Ancestors verified", "Warning", "Link"]


def rows(suggestions: list[dict]) -> list[list]:
    out = []
    for x in suggestions:
        lot = x["lot"]
        d = lot.score_detail or {}
        anc = d.get("ancestors") or {}
        out.append([
            lot.tier, x["rank_score"], lot.source, lot.auction_name, lot.lot_no, lot.ring_raw, lot.sex, lot.seller,
            x["price"], x["max"], fmt_local(lot.ends_at), lot.match_confidence,
            ", ".join(t["name"] for t in lot.family_tags or []),
            f'{anc.get("verified", 0)}/{anc.get("verified", 0) + anc.get("claimed", 0) + anc.get("mismatch", 0)}',
            "claim mismatch" if lot.claim_mismatch else ("seller-claimed, unverified" if lot.tier == "C" else ""),
            lot.url,
        ])
    return out


def to_csv(suggestions: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    w.writerows(rows(suggestions))
    return buf.getvalue()


def to_xlsx(suggestions: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Suggestions"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows(suggestions):
        ws.append(r)
    for col, width in zip("ABCDEFGHIJKLMNOP", [5, 7, 11, 28, 5, 18, 6, 22, 10, 10, 16, 9, 22, 10, 22, 50]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=10):
        for c in row:
            c.number_format = '"R"#,##0'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_whatsapp(suggestions: list[dict], limit: int = 25) -> str:
    """Monospace block (``` … ```) that WhatsApp renders as a table."""
    lines = ["*OLR Scout suggestions*", "```"]
    lines.append(f"{'T':1} {'Scr':>4} {'Lot':>4} {'Ring':<17} {'Now':>8} {'Max':>8} Closes")
    for x in suggestions[:limit]:
        lot = x["lot"]
        score = f"{x['rank_score']:.0f}" if x["rank_score"] is not None else "–"
        lines.append(
            f"{lot.tier or '-':1} {score:>4} {str(lot.lot_no or ''):>4} {(lot.ring_raw or '')[:17]:<17} "
            f"{format_zar(x['price']):>8} {format_zar(x['max']):>8} {fmt_local(lot.ends_at)}"
        )
    lines.append("```")
    tier_c = sum(1 for x in suggestions[:limit] if x["lot"].tier == "C")
    if tier_c:
        lines.append(f"_{tier_c} Tier C lot(s): seller-claimed, unverified._")
    return "\n".join(lines)
